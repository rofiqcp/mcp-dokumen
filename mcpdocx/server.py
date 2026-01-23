"""
MCP DOCX Unified Server - 90 Tools
A comprehensive MCP server for reading, generating, and editing Word documents.
"""

import argparse
import logging
from typing import Optional, List, Dict

from mcp.server.fastmcp import FastMCP

from mcpdocx.processor import DocumentProcessor, track_tool_usage

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


def create_server() -> tuple[FastMCP, DocumentProcessor]:
    """Create and configure the MCP server with 90 tools."""
    
    server = FastMCP("mcpdocx-unified")
    processor = DocumentProcessor()
    
    # =================================================================
    # DOCUMENT LIFECYCLE TOOLS (8 tools)
    # =================================================================
    
    # Tool 1
    @server.tool()
    async def create_document(file_path: str, title: Optional[str] = None) -> str:
        """Create a new Word document."""
        track_tool_usage("create_document")
        return processor.create_document(file_path, title)
    
    # Tool 2
    @server.tool()
    async def open_document(file_path: str) -> str:
        """Open an existing Word document for editing."""
        track_tool_usage("open_document")
        return processor.open_document(file_path)
    
    # Tool 3
    @server.tool()
    async def save_document(file_path: Optional[str] = None) -> str:
        """Save the current document. Optionally save to a new path."""
        track_tool_usage("save_document")
        return processor.save_document(file_path)
    
    # Tool 4
    @server.tool()
    async def close_document() -> str:
        """Close the current document without saving."""
        track_tool_usage("close_document")
        return processor.close_document()
    
    # Tool 5
    @server.tool()
    async def get_document_info() -> str:
        """Get information about the current document (sections, paragraphs, tables, word count)."""
        track_tool_usage("get_document_info")
        return processor.get_document_info()
    
    # Tool 6
    @server.tool()
    async def get_document_text() -> str:
        """Get all text content from the current document."""
        track_tool_usage("get_document_text")
        return processor.get_document_text()
    
    # Tool 7
    @server.tool()
    async def list_open_documents() -> str:
        """List all currently open documents."""
        track_tool_usage("list_open_documents")
        return processor.list_open_documents()
    
    # Tool 8
    @server.tool()
    async def switch_document(file_path: str) -> str:
        """Switch to a different open document."""
        track_tool_usage("switch_document")
        return processor.switch_document(file_path)
    
    # =================================================================
    # PARAGRAPH & CONTENT TOOLS (12 tools)
    # =================================================================
    
    # Tool 9
    @server.tool()
    async def add_paragraph(
        text: str,
        style: Optional[str] = None,
        bold: bool = False,
        italic: bool = False,
        underline: bool = False,
        font_size: Optional[int] = None,
        font_name: Optional[str] = None,
        color: Optional[str] = None,
        alignment: Optional[str] = None
    ) -> str:
        """Add a paragraph with optional formatting (bold, italic, underline, font, color, alignment)."""
        track_tool_usage("add_paragraph")
        return processor.add_paragraph(text, style, bold, italic, underline, font_size, font_name, color, alignment)
    
    # Tool 10
    @server.tool()
    async def add_heading(text: str, level: int = 1) -> str:
        """Add a heading to the document (level 0-9, where 0 is Title)."""
        track_tool_usage("add_heading")
        return processor.add_heading(text, level)
    
    # Tool 11
    @server.tool()
    async def add_page_break() -> str:
        """Add a page break to the document."""
        track_tool_usage("add_page_break")
        return processor.add_page_break()
    
    # Tool 12
    @server.tool()
    async def get_paragraph_text(index: int) -> str:
        """Get text of a specific paragraph by index (0-based)."""
        track_tool_usage("get_paragraph_text")
        return processor.get_paragraph_text(index)
    
    # Tool 13
    @server.tool()
    async def set_paragraph_text(index: int, text: str) -> str:
        """Replace the text of a specific paragraph by index."""
        track_tool_usage("set_paragraph_text")
        return processor.set_paragraph_text(index, text)
    
    # Tool 14
    @server.tool()
    async def delete_paragraph(index: int) -> str:
        """Delete a paragraph by index (0-based)."""
        track_tool_usage("delete_paragraph")
        return processor.delete_paragraph(index)
    
    # Tool 15
    @server.tool()
    async def insert_paragraph_after(index: int, text: str, style: Optional[str] = None) -> str:
        """Insert a new paragraph after the specified index."""
        track_tool_usage("insert_paragraph_after")
        return processor.insert_paragraph_after(index, text, style)
    
    # Tool 16
    @server.tool()
    async def count_paragraphs() -> str:
        """Get the total number of paragraphs in the document."""
        track_tool_usage("count_paragraphs")
        return processor.count_paragraphs()
    
    # Tool 17
    @server.tool()
    async def add_bulleted_list(items: List[str], style: str = "List Bullet") -> str:
        """Add a bulleted list to the document."""
        track_tool_usage("add_bulleted_list")
        return processor.add_bulleted_list(items, style)
    
    # Tool 18
    @server.tool()
    async def add_numbered_list(items: List[str], style: str = "List Number") -> str:
        """Add a numbered list to the document."""
        track_tool_usage("add_numbered_list")
        return processor.add_numbered_list(items, style)
    
    # Tool 19
    @server.tool()
    async def apply_style(paragraph_index: int, style_name: str) -> str:
        """Apply a style to a specific paragraph."""
        track_tool_usage("apply_style")
        return processor.apply_style(paragraph_index, style_name)
    
    # Tool 20
    @server.tool()
    async def batch_format_paragraphs(style_name: str, start_index: int = 0, end_index: Optional[int] = None) -> str:
        """Apply a style to a range of paragraphs."""
        track_tool_usage("batch_format_paragraphs")
        return processor.batch_format_paragraphs(style_name, start_index, end_index)
    
    # =================================================================
    # SEARCH & REPLACE TOOLS (5 tools)
    # =================================================================
    
    # Tool 21
    @server.tool()
    async def search_text(keyword: str) -> str:
        """Search for text in paragraphs and tables."""
        track_tool_usage("search_text")
        return processor.search_text(keyword)
    
    # Tool 22
    @server.tool()
    async def find_and_replace(find_text: str, replace_text: str) -> str:
        """Find and replace text throughout the document."""
        track_tool_usage("find_and_replace")
        return processor.find_and_replace(find_text, replace_text)
    
    # Tool 23
    @server.tool()
    async def batch_replace(replacements: Dict[str, str]) -> str:
        """Replace multiple strings at once. Pass dict of old->new text."""
        track_tool_usage("batch_replace")
        return processor.batch_replace(replacements)
    
    # Tool 24
    @server.tool()
    async def redact_text(pattern: str, replacement: str = "[REDACTED]", use_regex: bool = False) -> str:
        """Redact text matching pattern throughout the document."""
        track_tool_usage("redact_text")
        return processor.redact_text(pattern, replacement, use_regex)
    
    # Tool 25
    @server.tool()
    async def sanitize_external_links() -> str:
        """Remove all external hyperlinks from the document."""
        # This is in processor but not shown, add inline
        if not processor.current_document:
            return "Error: No document is open"
        try:
            from docx.oxml.ns import qn
            body = processor.current_document.element.body
            count = 0
            for hyperlink in list(body.findall('.//' + qn('w:hyperlink'))):
                r_id = hyperlink.get(qn('r:id'))
                if r_id:
                    parent = hyperlink.getparent()
                    index = list(parent).index(hyperlink)
                    for child in hyperlink:
                        parent.insert(index, child)
                        index += 1
                    parent.remove(hyperlink)
                    count += 1
            return f"Removed {count} external hyperlinks"
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # TABLE TOOLS (14 tools)
    # =================================================================
    
    # Tool 26
    @server.tool()
    async def add_table(rows: int, cols: int, data: Optional[List[List[str]]] = None, style: str = "Table Grid") -> str:
        """Add a table with optional initial data."""
        track_tool_usage("add_table")
        return processor.add_table(rows, cols, data, style)
    
    # Tool 27
    @server.tool()
    async def get_tables_info() -> str:
        """Get information about all tables in the document."""
        track_tool_usage("get_tables_info")
        return processor.get_tables_info()
    
    # Tool 28
    @server.tool()
    async def get_table_data(table_index: int) -> str:
        """Get all content from a specific table."""
        track_tool_usage("get_table_data")
        return processor.get_table_data(table_index)
    
    # Tool 29
    @server.tool()
    async def edit_table_cell(table_index: int, row_index: int, col_index: int, text: str) -> str:
        """Edit the content of a table cell."""
        track_tool_usage("edit_table_cell")
        return processor.edit_table_cell(table_index, row_index, col_index, text)
    
    # Tool 30
    @server.tool()
    async def add_table_row(table_index: int, data: Optional[List[str]] = None) -> str:
        """Add a row to an existing table with optional cell values."""
        track_tool_usage("add_table_row")
        return processor.add_table_row(table_index, data)
    
    # Tool 31
    @server.tool()
    async def delete_table_row(table_index: int, row_index: int) -> str:
        """Delete a row from a table."""
        track_tool_usage("delete_table_row")
        return processor.delete_table_row(table_index, row_index)
    
    # Tool 32
    @server.tool()
    async def merge_table_cells(table_index: int, start_row: int, start_col: int, end_row: int, end_col: int) -> str:
        """Merge cells in a rectangular area."""
        track_tool_usage("merge_table_cells")
        return processor.merge_table_cells(table_index, start_row, start_col, end_row, end_col)
    
    # Tool 33
    @server.tool()
    async def set_table_cell_shading(table_index: int, row_index: int, col_index: int, fill_color: str = "FFFF00") -> str:
        """Set cell background color (hex without #)."""
        track_tool_usage("set_table_cell_shading")
        return processor.set_table_cell_shading(table_index, row_index, col_index, fill_color)
    
    # Tool 34
    @server.tool()
    async def apply_table_alternating_rows(table_index: int, color1: str = "FFFFFF", color2: str = "F2F2F2") -> str:
        """Apply alternating row colors to a table."""
        track_tool_usage("apply_table_alternating_rows")
        return processor.apply_table_alternating_rows(table_index, color1, color2)
    
    # Tool 35
    @server.tool()
    async def highlight_table_header(table_index: int, header_color: str = "4472C4", text_color: str = "FFFFFF") -> str:
        """Highlight the first row as a table header."""
        track_tool_usage("highlight_table_header")
        return processor.highlight_table_header(table_index, header_color, text_color)
    
    # Tool 36
    @server.tool()
    async def set_column_width(table_index: int, col_index: int, width_cm: float) -> str:
        """Set width for a specific column in centimeters."""
        track_tool_usage("set_column_width")
        return processor.set_column_width(table_index, col_index, width_cm)
    
    # Tool 37
    @server.tool()
    async def set_row_height(table_index: int, row_index: int, height_cm: float, rule: str = "exact") -> str:
        """Set height for a specific row in centimeters."""
        track_tool_usage("set_row_height")
        return processor.set_row_height(table_index, row_index, height_cm, rule)
    
    # Tool 38
    @server.tool()
    async def set_table_borders(
        table_index: int,
        border_style: str = "single",
        size: int = 4,
        color: str = "000000"
    ) -> str:
        """Set table border style (single, double, dotted, dashed, none)."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            tables = processor.current_document.tables
            if table_index >= len(tables):
                return "Error: Invalid table index"
            table = tables[table_index]
            tbl = table._tbl
            tblPr = tbl.tblPr
            if tblPr is None:
                tblPr = OxmlElement('w:tblPr')
                tbl.insert(0, tblPr)
            tblBorders = tblPr.find(qn('w:tblBorders'))
            if tblBorders is None:
                tblBorders = OxmlElement('w:tblBorders')
                tblPr.append(tblBorders)
            for name in ['top', 'bottom', 'left', 'right', 'insideH', 'insideV']:
                existing = tblBorders.find(qn(f'w:{name}'))
                if existing is not None:
                    tblBorders.remove(existing)
                border = OxmlElement(f'w:{name}')
                border.set(qn('w:val'), border_style if border_style != 'none' else 'nil')
                border.set(qn('w:sz'), str(size))
                border.set(qn('w:color'), color)
                tblBorders.append(border)
            return "Table borders updated"
        except Exception as e:
            return f"Error: {e}"
    
    # Tool 39
    @server.tool()
    async def import_csv_as_table(csv_path: str, delimiter: str = ",") -> str:
        """Import a CSV file as a table in the document."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            import csv
            import os
            if not os.path.exists(csv_path):
                return f"Error: File not found: {csv_path}"
            with open(csv_path, 'r', newline='', encoding='utf-8') as f:
                reader = csv.reader(f, delimiter=delimiter)
                rows = list(reader)
            if not rows:
                return "Error: Empty CSV file"
            num_rows = len(rows)
            num_cols = max(len(row) for row in rows)
            table = processor.current_document.add_table(rows=num_rows, cols=num_cols)
            table.style = 'Table Grid'
            for r_idx, row_data in enumerate(rows):
                for c_idx, value in enumerate(row_data):
                    if c_idx < num_cols:
                        table.cell(r_idx, c_idx).text = str(value)
            return f"Added table with {num_rows} rows and {num_cols} columns from CSV"
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # HEADER & FOOTER TOOLS (8 tools)
    # =================================================================
    
    # Tool 40
    @server.tool()
    async def add_header(text: str, section_index: int = 0, alignment: str = "center") -> str:
        """Add or update header text for a section."""
        track_tool_usage("add_header")
        return processor.add_header(text, section_index, alignment)
    
    # Tool 41
    @server.tool()
    async def add_footer(text: str, section_index: int = 0, alignment: str = "center") -> str:
        """Add or update footer text for a section."""
        track_tool_usage("add_footer")
        return processor.add_footer(text, section_index, alignment)
    
    # Tool 42
    @server.tool()
    async def add_page_numbers(position: str = "footer", alignment: str = "center", format_string: str = "Page {page}", section_index: int = 0) -> str:
        """Add page numbers to header or footer."""
        track_tool_usage("add_page_numbers")
        return processor.add_page_numbers(position, alignment, format_string, section_index)
    
    # Tool 43
    @server.tool()
    async def get_header_text(section_index: int = 0) -> str:
        """Get the header text from a section."""
        track_tool_usage("get_header_text")
        return processor.get_header_text(section_index)
    
    # Tool 44
    @server.tool()
    async def get_footer_text(section_index: int = 0) -> str:
        """Get the footer text from a section."""
        track_tool_usage("get_footer_text")
        return processor.get_footer_text(section_index)
    
    # Tool 45
    @server.tool()
    async def remove_header(section_index: int = 0) -> str:
        """Remove header from a section."""
        track_tool_usage("remove_header")
        return processor.remove_header(section_index)
    
    # Tool 46
    @server.tool()
    async def remove_footer(section_index: int = 0) -> str:
        """Remove footer from a section."""
        track_tool_usage("remove_footer")
        return processor.remove_footer(section_index)
    
    # Tool 47
    @server.tool()
    async def add_text_watermark(text: str = "CONFIDENTIAL", section_index: int = 0) -> str:
        """Add a text watermark via header."""
        track_tool_usage("add_text_watermark")
        return processor.add_text_watermark(text, section_index)
    
    # =================================================================
    # IMAGE TOOLS (4 tools)
    # =================================================================
    
    # Tool 48
    @server.tool()
    async def add_image(image_path: str, width_inches: Optional[float] = None, height_inches: Optional[float] = None) -> str:
        """Add an image to the document with optional size."""
        track_tool_usage("add_image")
        return processor.add_image(image_path, width_inches, height_inches)
    
    # Tool 49
    @server.tool()
    async def add_image_base64(base64_data: str, width_inches: Optional[float] = None) -> str:
        """Add an image from base64 encoded data."""
        track_tool_usage("add_image_base64")
        return processor.add_image_base64(base64_data, width_inches)
    
    # Tool 50
    @server.tool()
    async def list_images() -> str:
        """List all images in the document with dimensions."""
        track_tool_usage("list_images")
        return processor.list_images()
    
    # Tool 51
    @server.tool()
    async def add_image_watermark(image_path: str, width_inches: float = 4.0, section_index: int = 0) -> str:
        """Add an image as a watermark via header."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            import os
            from docx.shared import Inches
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
            if not os.path.exists(image_path):
                return f"Error: Image not found: {image_path}"
            sections = processor.current_document.sections
            if section_index >= len(sections):
                return f"Error: Section index out of range"
            header = sections[section_index].header
            para = header.add_paragraph()
            para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = para.add_run()
            run.add_picture(image_path, width=Inches(width_inches))
            return f"Image watermark added"
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # PAGE LAYOUT TOOLS (8 tools)
    # =================================================================
    
    # Tool 52
    @server.tool()
    async def set_page_margins(top: Optional[float] = None, bottom: Optional[float] = None, left: Optional[float] = None, right: Optional[float] = None, section_index: int = 0) -> str:
        """Set page margins in centimeters."""
        track_tool_usage("set_page_margins")
        return processor.set_page_margins(top, bottom, left, right, section_index)
    
    # Tool 53
    @server.tool()
    async def set_page_size(width: float, height: float, section_index: int = 0) -> str:
        """Set page size in centimeters."""
        track_tool_usage("set_page_size")
        return processor.set_page_size(width, height, section_index)
    
    # Tool 54
    @server.tool()
    async def set_page_orientation(orientation: str, section_index: int = 0) -> str:
        """Set page orientation ('portrait' or 'landscape')."""
        track_tool_usage("set_page_orientation")
        return processor.set_page_orientation(orientation, section_index)
    
    # Tool 55
    @server.tool()
    async def get_page_info(section_index: int = 0) -> str:
        """Get page layout information (size, orientation, margins)."""
        track_tool_usage("get_page_info")
        return processor.get_page_info(section_index)
    
    # Tool 56
    @server.tool()
    async def add_section(section_type: str = "NEW_PAGE") -> str:
        """Add a new section (NEW_PAGE, CONTINUOUS, ODD_PAGE, EVEN_PAGE)."""
        track_tool_usage("add_section")
        return processor.add_section(section_type)
    
    # Tool 57
    @server.tool()
    async def list_sections() -> str:
        """List all sections with their properties."""
        track_tool_usage("list_sections")
        return processor.list_sections()
    
    # Tool 58
    @server.tool()
    async def list_styles() -> str:
        """List available paragraph styles in the document."""
        track_tool_usage("list_styles")
        return processor.list_styles()
    
    # Tool 59
    @server.tool()
    async def get_formatting_report() -> str:
        """Get a report of fonts and styles used in the document."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            fonts_used = set()
            sizes_used = set()
            styles_used = set()
            for para in processor.current_document.paragraphs:
                if para.style:
                    styles_used.add(para.style.name)
                for run in para.runs:
                    if run.font.name:
                        fonts_used.add(run.font.name)
                    if run.font.size:
                        sizes_used.add(str(run.font.size.pt) + "pt")
            return f"""Formatting Report:- Fonts: {', '.join(sorted(fonts_used)) or 'Default only'}
- Sizes: {', '.join(sorted(sizes_used)) or 'Default only'}
- Styles: {', '.join(sorted(styles_used)) or 'None'}"""
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # METADATA TOOLS (3 tools)
    # =================================================================
    
    # Tool 60
    @server.tool()
    async def get_metadata() -> str:
        """Get document metadata (title, author, subject, keywords, dates)."""
        track_tool_usage("get_metadata")
        return processor.get_metadata()
    
    # Tool 61
    @server.tool()
    async def set_metadata(title: Optional[str] = None, author: Optional[str] = None, subject: Optional[str] = None, keywords: Optional[str] = None) -> str:
        """Set document metadata properties."""
        track_tool_usage("set_metadata")
        return processor.set_metadata(title, author, subject, keywords)
    
    # Tool 62
    @server.tool()
    async def strip_personal_info() -> str:
        """Remove personal information from document metadata."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            props = processor.current_document.core_properties
            props.author = ""
            props.last_modified_by = ""
            props.comments = ""
            return "Personal information stripped from metadata"
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # COMMENT & ANNOTATION TOOLS (4 tools)
    # =================================================================
    
    # Tool 63
    @server.tool()
    async def add_comment(comment_text: str, paragraph_index: int, author: str = "MCP") -> str:
        """Add a comment marker to a paragraph."""
        track_tool_usage("add_comment")
        return processor.add_comment(comment_text, paragraph_index, author)
    
    # Tool 64
    @server.tool()
    async def get_all_comments() -> str:
        """Get all comments from the document."""
        track_tool_usage("get_all_comments")
        return processor.get_all_comments()
    
    # Tool 65
    @server.tool()
    async def delete_all_comments() -> str:
        """Delete all comments from the document."""
        track_tool_usage("delete_all_comments")
        return processor.delete_all_comments()
    
    # Tool 66
    @server.tool()
    async def add_footnote(text: str, paragraph_index: int) -> str:
        """Add a footnote to a paragraph."""
        track_tool_usage("add_footnote")
        return processor.add_footnote(text, paragraph_index)
    
    # =================================================================
    # HYPERLINK & BOOKMARK TOOLS (4 tools)
    # =================================================================
    
    # Tool 67
    @server.tool()
    async def add_hyperlink(paragraph_index: int, text: str, url: str) -> str:
        """Add an external hyperlink to a paragraph."""
        track_tool_usage("add_hyperlink")
        return processor.add_hyperlink(paragraph_index, text, url)
    
    # Tool 68
    @server.tool()
    async def add_bookmark(paragraph_index: int, bookmark_name: str) -> str:
        """Add a bookmark to a paragraph for internal navigation."""
        track_tool_usage("add_bookmark")
        return processor.add_bookmark(paragraph_index, bookmark_name)
    
    # Tool 69
    @server.tool()
    async def list_bookmarks() -> str:
        """List all bookmarks in the document."""
        track_tool_usage("list_bookmarks")
        return processor.list_bookmarks()
    
    # Tool 70
    @server.tool()
    async def add_internal_link(text: str, bookmark_name: str) -> str:
        """Add a hyperlink that navigates to a bookmark."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            para = processor.current_document.add_paragraph()
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('w:anchor'), bookmark_name)
            new_run = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')
            color = OxmlElement('w:color')
            color.set(qn('w:val'), '0000FF')
            rPr.append(color)
            u = OxmlElement('w:u')
            u.set(qn('w:val'), 'single')
            rPr.append(u)
            new_run.append(rPr)
            t = OxmlElement('w:t')
            t.text = text
            new_run.append(t)
            hyperlink.append(new_run)
            para._p.append(hyperlink)
            return f"Internal link to bookmark '{bookmark_name}' added"
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # DOCUMENT MERGE & COPY TOOLS (3 tools)
    # =================================================================
    
    # Tool 71
    @server.tool()
    async def merge_documents(file_paths: List[str]) -> str:
        """Merge multiple documents into the current document."""
        track_tool_usage("merge_documents")
        return processor.merge_documents(file_paths)
    
    # Tool 72
    @server.tool()
    async def copy_document(dest_path: str) -> str:
        """Create a copy of the current document."""
        track_tool_usage("copy_document")
        return processor.copy_document(dest_path)
    
    # Tool 73
    @server.tool()
    async def insert_document(file_path: str) -> str:
        """Insert another document at the end of the current document."""
        track_tool_usage("insert_document")
        return processor.merge_documents([file_path])
    
    # =================================================================
    # TOC & STRUCTURE TOOLS (4 tools)
    # =================================================================
    
    # Tool 74
    @server.tool()
    async def add_table_of_contents(title: str = "Table of Contents", heading_levels: int = 3) -> str:
        """Add a Table of Contents to the document."""
        track_tool_usage("add_table_of_contents")
        return processor.add_table_of_contents(title, heading_levels)
    
    # Tool 75
    @server.tool()
    async def get_document_outline() -> str:
        """Get document outline based on headings."""
        track_tool_usage("get_document_outline")
        return processor.get_document_outline()
    
    # Tool 76
    @server.tool()
    async def get_document_structure() -> str:
        """Get structural overview (sections, paragraphs, tables, headings)."""
        track_tool_usage("get_document_structure")
        return processor.get_document_structure()
    
    # Tool 77
    @server.tool()
    async def get_raw_xml() -> str:
        """Get the raw XML structure of the document body (for debugging)."""
        if not processor.current_document:
            return "Error: No document is open"
        try:
            from lxml import etree
            body = processor.current_document.element.body
            xml_str = etree.tostring(body, pretty_print=True, encoding='unicode')
            if len(xml_str) > 10000:
                return xml_str[:10000] + "\n... (truncated)"
            return xml_str
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # TRACK CHANGES & PROTECTION TOOLS (6 tools)
    # =================================================================
    
    # Tool 78
    @server.tool()
    async def enable_track_changes() -> str:
        """Enable track changes in the document."""
        track_tool_usage("enable_track_changes")
        return processor.enable_track_changes()
    
    # Tool 79
    @server.tool()
    async def disable_track_changes() -> str:
        """Disable track changes in the document."""
        track_tool_usage("disable_track_changes")
        return processor.disable_track_changes()
    
    # Tool 80
    @server.tool()
    async def accept_all_changes() -> str:
        """Accept all tracked changes in the document."""
        track_tool_usage("accept_all_changes")
        return processor.accept_all_changes()
    
    # Tool 81
    @server.tool()
    async def reject_all_changes() -> str:
        """Reject all tracked changes in the document."""
        track_tool_usage("reject_all_changes")
        return processor.reject_all_changes()
    
    # Tool 82
    @server.tool()
    async def protect_document(password: Optional[str] = None, protection_type: str = "readOnly") -> str:
        """Protect document (readOnly, comments, trackedChanges, forms)."""
        track_tool_usage("protect_document")
        return processor.protect_document(password, protection_type)
    
    # Tool 83
    @server.tool()
    async def unprotect_document() -> str:
        """Remove document protection."""
        track_tool_usage("unprotect_document")
        return processor.unprotect_document()
    
    # =================================================================
    # STATISTICS TOOLS (3 tools)
    # =================================================================
    
    # Tool 84
    @server.tool()
    async def get_word_count() -> str:
        """Get word count statistics (words, characters, paragraphs)."""
        track_tool_usage("get_word_count")
        return processor.get_word_count()
    
    # Tool 85
    @server.tool()
    async def get_document_statistics() -> str:
        """Get comprehensive document statistics."""
        track_tool_usage("get_document_statistics")
        return processor.get_document_statistics()
    
    # Tool 86
    @server.tool()
    async def count_tables() -> str:
        """Get the number of tables in the document."""
        if not processor.current_document:
            return "Error: No document is open"
        return f"Document has {len(processor.current_document.tables)} tables"
    
    # =================================================================
    # FILE OPERATION TOOLS (2 tools)
    # =================================================================
    
    # Tool 87
    @server.tool()
    async def list_docx_files(directory: str) -> str:
        """List all .docx files in a directory."""
        track_tool_usage("list_docx_files")
        return processor.list_docx_files(directory)
    
    # Tool 88
    @server.tool()
    async def compare_documents(file_path1: str, file_path2: str) -> str:
        """Compare two documents and show differences in paragraph count and word count."""
        try:
            from docx import Document
            import os
            if not os.path.exists(file_path1):
                return f"Error: File not found: {file_path1}"
            if not os.path.exists(file_path2):
                return f"Error: File not found: {file_path2}"
            doc1 = Document(file_path1)
            doc2 = Document(file_path2)
            para1 = len(doc1.paragraphs)
            para2 = len(doc2.paragraphs)
            words1 = sum(len(p.text.split()) for p in doc1.paragraphs)
            words2 = sum(len(p.text.split()) for p in doc2.paragraphs)
            tables1 = len(doc1.tables)
            tables2 = len(doc2.tables)
            return f"""Document Comparison:File 1: {file_path1}
  - Paragraphs: {para1}
  - Words: {words1}
  - Tables: {tables1}

File 2: {file_path2}
  - Paragraphs: {para2}
  - Words: {words2}
  - Tables: {tables2}

Differences:
  - Paragraphs: {para2 - para1:+d}
  - Words: {words2 - words1:+d}
  - Tables: {tables2 - tables1:+d}"""
        except Exception as e:
            return f"Error: {e}"
    
    # =================================================================
    # MAIL MERGE TOOLS (2 tools)
    # =================================================================
    
    # Tool 89
    @server.tool()
    async def add_merge_field(field_name: str) -> str:
        """Add a mail merge field placeholder to the document."""
        track_tool_usage("add_merge_field")
        return processor.add_merge_field(field_name)
    
    # Tool 90
    @server.tool()
    async def execute_mail_merge(data: List[Dict[str, str]], output_pattern: str) -> str:
        """Execute mail merge with data, creating one document per record. Use {index} in output_pattern."""
        track_tool_usage("execute_mail_merge")
        return processor.execute_mail_merge(data, output_pattern)
    
    # =================================================================
    
    logger.info("MCP DOCX Unified server initialized with 90 tools")
    
    return server, processor


def run_server():
    """Run the MCP server using stdio transport."""
    server, _ = create_server()
    logger.info("Starting MCP DOCX Unified server...")
    server.run()


def main():
    """Main entry point."""
    parser = argparse.ArgumentParser(
        description="MCP DOCX Unified - 90 Tools for Word Document Operations"
    )
    parser.add_argument(
        "--version",
        action="version",
        version="MCP DOCX Unified 1.0.0"
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Enable debug logging"
    )
    
    args = parser.parse_args()
    
    if args.debug:
        logging.getLogger().setLevel(logging.DEBUG)
    
    run_server()


if __name__ == "__main__":
    main()
