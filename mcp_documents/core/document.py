"""
Core document processor class
"""

import os
import logging
from typing import Optional, Dict, Any, List

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from mcp_documents.core.formatting import apply_paragraph_formatting, apply_run_formatting
from mcp_documents.core.utils import parse_color, logger


class DocumentProcessor:
    """
    Unified document processor for Word documents.
    Manages document lifecycle and provides all editing operations.
    """
    
    def __init__(self):
        self.documents: Dict[str, Document] = {}  # Cache of opened documents
        self.current_document: Optional[Document] = None
        self.current_file_path: Optional[str] = None
    
    # ==================== Document Lifecycle ====================
    
    def create_document(self, file_path: str, title: Optional[str] = None) -> str:
        """Create a new Word document."""
        try:
            doc = Document()
            
            if title:
                doc.add_heading(title, 0)
            
            doc.save(file_path)
            
            self.current_document = doc
            self.current_file_path = file_path
            self.documents[file_path] = doc
            
            return f"Document created: {file_path}"
        except Exception as e:
            logger.error(f"Failed to create document: {e}")
            return f"Error: {e}"
    
    def open_document(self, file_path: str) -> str:
        """Open an existing Word document."""
        try:
            if not os.path.exists(file_path):
                return f"Error: File not found: {file_path}"
            
            doc = Document(file_path)
            
            self.current_document = doc
            self.current_file_path = file_path
            self.documents[file_path] = doc
            
            return f"Document opened: {file_path}"
        except Exception as e:
            logger.error(f"Failed to open document: {e}")
            return f"Error: {e}"
    
    def save_document(self, file_path: Optional[str] = None) -> str:
        """Save the current document."""
        try:
            if not self.current_document:
                return "Error: No document is open"
            
            save_path = file_path or self.current_file_path
            if not save_path:
                return "Error: No file path specified"
            
            self.current_document.save(save_path)
            
            if file_path:
                self.current_file_path = file_path
                self.documents[file_path] = self.current_document
            
            return f"Document saved: {save_path}"
        except Exception as e:
            logger.error(f"Failed to save document: {e}")
            return f"Error: {e}"
    
    def close_document(self) -> str:
        """Close the current document."""
        if self.current_file_path and self.current_file_path in self.documents:
            del self.documents[self.current_file_path]
        
        self.current_document = None
        self.current_file_path = None
        
        return "Document closed"
    
    def get_document_info(self) -> str:
        """Get information about the current document."""
        if not self.current_document:
            return "Error: No document is open"
        
        doc = self.current_document
        
        info = []
        info.append(f"File: {self.current_file_path}")
        info.append(f"Sections: {len(doc.sections)}")
        info.append(f"Paragraphs: {len(doc.paragraphs)}")
        info.append(f"Tables: {len(doc.tables)}")
        
        # Word count
        word_count = sum(len(p.text.split()) for p in doc.paragraphs)
        info.append(f"Word count: {word_count}")
        
        return "\n".join(info)
    
    # ==================== Content Operations ====================
    
    def add_paragraph(
        self,
        text: str,
        style: Optional[str] = None,
        bold: bool = False,
        italic: bool = False,
        underline: bool = False,
        font_size: Optional[int] = None,
        font_name: Optional[str] = None,
        color: Optional[str] = None,
        alignment: Optional[str] = None
    ) -> str:
        """Add a paragraph with optional formatting."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            # Add paragraph with optional style
            if style:
                try:
                    paragraph = self.current_document.add_paragraph(text, style=style)
                except KeyError:
                    paragraph = self.current_document.add_paragraph(text)
            else:
                paragraph = self.current_document.add_paragraph(text)
            
            # Apply text formatting to runs
            if paragraph.runs:
                run = paragraph.runs[0]
                run.bold = bold
                run.italic = italic
                run.underline = underline
                
                if font_size:
                    run.font.size = Pt(font_size)
                
                if font_name:
                    run.font.name = font_name
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                
                if color:
                    rgb = parse_color(color)
                    if rgb:
                        run.font.color.rgb = RGBColor(*rgb)
            
            # Apply paragraph alignment
            if alignment:
                alignment_map = {
                    "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
                    "center": WD_PARAGRAPH_ALIGNMENT.CENTER,
                    "right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
                    "justify": WD_PARAGRAPH_ALIGNMENT.JUSTIFY
                }
                if alignment.lower() in alignment_map:
                    paragraph.alignment = alignment_map[alignment.lower()]
            
            return "Paragraph added"
        except Exception as e:
            logger.error(f"Failed to add paragraph: {e}")
            return f"Error: {e}"
    
    def add_heading(self, text: str, level: int = 1) -> str:
        """Add a heading to the document."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            self.current_document.add_heading(text, level=level)
            return f"Heading level {level} added"
        except Exception as e:
            logger.error(f"Failed to add heading: {e}")
            return f"Error: {e}"
    
    def add_table(
        self,
        rows: int,
        cols: int,
        data: Optional[List[List[str]]] = None,
        style: str = "Table Grid"
    ) -> str:
        """Add a table to the document."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            table = self.current_document.add_table(rows=rows, cols=cols, style=style)
            
            # Fill table data
            if data:
                for i, row_data in enumerate(data):
                    if i < rows:
                        row = table.rows[i]
                        for j, cell_text in enumerate(row_data):
                            if j < cols:
                                row.cells[j].text = str(cell_text)
            
            return f"Table {rows}x{cols} added"
        except Exception as e:
            logger.error(f"Failed to add table: {e}")
            return f"Error: {e}"
    
    def add_image(
        self,
        image_path: str,
        width_inches: Optional[float] = None,
        height_inches: Optional[float] = None
    ) -> str:
        """Add an image to the document."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            if not os.path.exists(image_path):
                return f"Error: Image not found: {image_path}"
            
            width = Inches(width_inches) if width_inches else None
            height = Inches(height_inches) if height_inches else None
            
            self.current_document.add_picture(image_path, width=width, height=height)
            
            return "Image added"
        except Exception as e:
            logger.error(f"Failed to add image: {e}")
            return f"Error: {e}"
    
    def add_page_break(self) -> str:
        """Add a page break."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            self.current_document.add_page_break()
            return "Page break added"
        except Exception as e:
            logger.error(f"Failed to add page break: {e}")
            return f"Error: {e}"
    
    # ==================== Table Operations ====================
    
    def edit_table_cell(
        self,
        table_index: int,
        row_index: int,
        col_index: int,
        text: str
    ) -> str:
        """Edit a table cell's content."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            tables = self.current_document.tables
            if table_index >= len(tables):
                return f"Error: Table index out of range (have {len(tables)} tables)"
            
            table = tables[table_index]
            if row_index >= len(table.rows):
                return f"Error: Row index out of range"
            if col_index >= len(table.columns):
                return f"Error: Column index out of range"
            
            table.cell(row_index, col_index).text = text
            return f"Cell ({row_index},{col_index}) updated"
        except Exception as e:
            logger.error(f"Failed to edit table cell: {e}")
            return f"Error: {e}"
    
    def add_table_row(self, table_index: int, data: Optional[List[str]] = None) -> str:
        """Add a row to a table."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            tables = self.current_document.tables
            if table_index >= len(tables):
                return f"Error: Table index out of range"
            
            table = tables[table_index]
            new_row = table.add_row()
            
            if data:
                for i, cell_text in enumerate(data):
                    if i < len(new_row.cells):
                        new_row.cells[i].text = str(cell_text)
            
            return "Row added"
        except Exception as e:
            logger.error(f"Failed to add table row: {e}")
            return f"Error: {e}"
    
    def delete_table_row(self, table_index: int, row_index: int) -> str:
        """Delete a row from a table."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            tables = self.current_document.tables
            if table_index >= len(tables):
                return f"Error: Table index out of range"
            
            table = tables[table_index]
            if row_index >= len(table.rows):
                return f"Error: Row index out of range"
            
            row = table.rows[row_index]._tr
            row.getparent().remove(row)
            
            return f"Row {row_index} deleted"
        except Exception as e:
            logger.error(f"Failed to delete table row: {e}")
            return f"Error: {e}"
    
    def merge_table_cells(
        self,
        table_index: int,
        start_row: int,
        start_col: int,
        end_row: int,
        end_col: int
    ) -> str:
        """Merge table cells."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            tables = self.current_document.tables
            if table_index >= len(tables):
                return f"Error: Table index out of range"
            
            table = tables[table_index]
            start_cell = table.cell(start_row, start_col)
            end_cell = table.cell(end_row, end_col)
            start_cell.merge(end_cell)
            
            return f"Cells merged from ({start_row},{start_col}) to ({end_row},{end_col})"
        except Exception as e:
            logger.error(f"Failed to merge cells: {e}")
            return f"Error: {e}"
    
    # ==================== Search and Edit ====================
    
    def search_text(self, keyword: str) -> str:
        """Search for text in the document."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            results = []
            
            # Search paragraphs
            for i, para in enumerate(self.current_document.paragraphs):
                if keyword in para.text:
                    results.append(f"Paragraph {i}: {para.text[:100]}...")
            
            # Search tables
            for t_idx, table in enumerate(self.current_document.tables):
                for r_idx, row in enumerate(table.rows):
                    for c_idx, cell in enumerate(row.cells):
                        if keyword in cell.text:
                            results.append(f"Table {t_idx} cell ({r_idx},{c_idx}): {cell.text[:50]}...")
            
            if not results:
                return f"'{keyword}' not found"
            
            return f"Found {len(results)} matches:\n" + "\n".join(results)
        except Exception as e:
            logger.error(f"Search failed: {e}")
            return f"Error: {e}"
    
    def find_and_replace(self, find_text: str, replace_text: str) -> str:
        """Find and replace text in the document."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            count = 0
            
            # Replace in paragraphs
            for para in self.current_document.paragraphs:
                if find_text in para.text:
                    occurrences = para.text.count(find_text)
                    para.text = para.text.replace(find_text, replace_text)
                    count += occurrences
            
            # Replace in tables
            for table in self.current_document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            if find_text in para.text:
                                occurrences = para.text.count(find_text)
                                para.text = para.text.replace(find_text, replace_text)
                                count += occurrences
            
            return f"Replaced {count} occurrences"
        except Exception as e:
            logger.error(f"Replace failed: {e}")
            return f"Error: {e}"
    
    def delete_paragraph(self, index: int) -> str:
        """Delete a paragraph by index."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            paragraphs = self.current_document.paragraphs
            if index >= len(paragraphs):
                return f"Error: Paragraph index out of range"
            
            p = paragraphs[index]._element
            p.getparent().remove(p)
            
            return f"Paragraph {index} deleted"
        except Exception as e:
            logger.error(f"Failed to delete paragraph: {e}")
            return f"Error: {e}"
    
    # ==================== Page Layout ====================
    
    def set_page_margins(
        self,
        top: Optional[float] = None,
        bottom: Optional[float] = None,
        left: Optional[float] = None,
        right: Optional[float] = None,
        section_index: int = 0
    ) -> str:
        """Set page margins in centimeters."""
        if not self.current_document:
            return "Error: No document is open"
        
        try:
            sections = self.current_document.sections
            if section_index >= len(sections):
                return f"Error: Section index out of range"
            
            section = sections[section_index]
            
            if top is not None:
                section.top_margin = Cm(top)
            if bottom is not None:
                section.bottom_margin = Cm(bottom)
            if left is not None:
                section.left_margin = Cm(left)
            if right is not None:
                section.right_margin = Cm(right)
            
            return "Page margins updated"
        except Exception as e:
            logger.error(f"Failed to set margins: {e}")
            return f"Error: {e}"
    
    # ==================== Paragraphs and Sections Access ====================
    
    def get_paragraphs(self) -> List[str]:
        """Get list of all paragraph texts."""
        if not self.current_document:
            return []
        return [p.text for p in self.current_document.paragraphs]
    
    def get_tables_info(self) -> str:
        """Get information about all tables."""
        if not self.current_document:
            return "Error: No document is open"
        
        tables = self.current_document.tables
        if not tables:
            return "No tables in document"
        
        info = []
        for i, table in enumerate(tables):
            rows = len(table.rows)
            cols = len(table.columns)
            info.append(f"Table {i}: {rows} rows x {cols} columns")
        
        return "\n".join(info)
