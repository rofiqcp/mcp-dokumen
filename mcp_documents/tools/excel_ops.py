"""
Excel/CSV operations for MCP Document Server
"""

import os
import csv
from typing import Optional, List, Dict, Any
import logging

logger = logging.getLogger(__name__)


def register_excel_tools(mcp_server, processor):
    """Register Excel/CSV tools with the MCP server."""

    @mcp_server.tool()
    def read_csv(file_path: str, delimiter: str = ",", 
                 has_header: bool = True) -> str:
        """
        Read a CSV file and return its contents.
        
        Args:
            file_path: Path to the CSV file
            delimiter: Column delimiter (default comma)
            has_header: Whether the first row is a header
        
        Returns:
            CSV contents as formatted text or error message
        """
        try:
            if not os.path.exists(file_path):
                return f"Error: File not found: {file_path}"
            
            with open(file_path, 'r', newline='', encoding='utf-8') as f:
                reader = csv.reader(f, delimiter=delimiter)
                rows = list(reader)
            
            if not rows:
                return "Empty CSV file"
            
            result = []
            if has_header:
                result.append(f"Headers: {', '.join(rows[0])}")
                result.append(f"Rows: {len(rows) - 1}")
                for i, row in enumerate(rows[1:6], 1):  # Show first 5 data rows
                    result.append(f"  Row {i}: {row}")
                if len(rows) > 6:
                    result.append(f"  ... and {len(rows) - 6} more rows")
            else:
                result.append(f"Rows: {len(rows)}")
                for i, row in enumerate(rows[:5], 1):
                    result.append(f"  Row {i}: {row}")
                if len(rows) > 5:
                    result.append(f"  ... and {len(rows) - 5} more rows")
            
            return "\n".join(result)
        except Exception as e:
            logger.error(f"Failed to read CSV: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def write_csv(file_path: str, data: List[List[str]], 
                  delimiter: str = ",") -> str:
        """
        Write data to a CSV file.
        
        Args:
            file_path: Path for the output CSV file
            data: List of rows, each row is a list of values
            delimiter: Column delimiter (default comma)
        
        Returns:
            Success or error message
        """
        try:
            with open(file_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f, delimiter=delimiter)
                writer.writerows(data)
            
            return f"CSV written to {file_path} ({len(data)} rows)"
        except Exception as e:
            logger.error(f"Failed to write CSV: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def csv_to_table(csv_path: str, delimiter: str = ",") -> str:
        """
        Import a CSV file as a table in the current document.
        
        Args:
            csv_path: Path to the CSV file
            delimiter: Column delimiter (default comma)
        
        Returns:
            Success or error message
        """
        try:
            if not processor.current_document:
                return "Error: No document is open"
            
            if not os.path.exists(csv_path):
                return f"Error: File not found: {csv_path}"
            
            with open(csv_path, 'r', newline='', encoding='utf-8') as f:
                reader = csv.reader(f, delimiter=delimiter)
                rows = list(reader)
            
            if not rows:
                return "Error: Empty CSV file"
            
            num_rows = len(rows)
            num_cols = max(len(row) for row in rows)
            
            table = processor.current_document.add_table(
                rows=num_rows, cols=num_cols
            )
            table.style = 'Table Grid'
            
            for r_idx, row_data in enumerate(rows):
                for c_idx, value in enumerate(row_data):
                    if c_idx < num_cols:
                        table.cell(r_idx, c_idx).text = str(value)
            
            return f"Added table with {num_rows} rows and {num_cols} columns from CSV"
        except Exception as e:
            logger.error(f"Failed to convert CSV to table: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def table_to_csv(table_index: int, output_path: str, 
                     delimiter: str = ",") -> str:
        """
        Export a table from the document to a CSV file.
        
        Args:
            table_index: Index of the table to export
            output_path: Path for the output CSV file
            delimiter: Column delimiter (default comma)
        
        Returns:
            Success or error message
        """
        try:
            if not processor.current_document:
                return "Error: No document is open"
            
            tables = processor.current_document.tables
            if table_index >= len(tables):
                return f"Error: Invalid table index"
            
            table = tables[table_index]
            rows = []
            
            for row in table.rows:
                row_data = [cell.text for cell in row.cells]
                rows.append(row_data)
            
            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f, delimiter=delimiter)
                writer.writerows(rows)
            
            return f"Exported table to {output_path} ({len(rows)} rows)"
        except Exception as e:
            logger.error(f"Failed to export table to CSV: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def read_excel(file_path: str, sheet_name: Optional[str] = None) -> str:
        """
        Read an Excel file and return its contents.
        Requires openpyxl package.
        
        Args:
            file_path: Path to the Excel file (.xlsx)
            sheet_name: Name of the sheet to read (default: active sheet)
        
        Returns:
            Excel contents as formatted text or error message
        """
        try:
            import openpyxl
        except ImportError:
            return "Error: openpyxl package required. Install with: pip install openpyxl"
        
        try:
            if not os.path.exists(file_path):
                return f"Error: File not found: {file_path}"
            
            wb = openpyxl.load_workbook(file_path)
            
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    return f"Error: Sheet '{sheet_name}' not found. Available: {', '.join(wb.sheetnames)}"
                ws = wb[sheet_name]
            else:
                ws = wb.active
            
            result = [f"Sheet: {ws.title}", f"Dimensions: {ws.dimensions}"]
            
            for i, row in enumerate(ws.iter_rows(max_row=6, values_only=True), 1):
                result.append(f"  Row {i}: {list(row)}")
            
            if ws.max_row > 6:
                result.append(f"  ... and {ws.max_row - 6} more rows")
            
            wb.close()
            return "\n".join(result)
        except Exception as e:
            logger.error(f"Failed to read Excel: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def excel_to_csv(excel_path: str, csv_path: str, 
                     sheet_name: Optional[str] = None) -> str:
        """
        Convert an Excel file to CSV format.
        Requires openpyxl package.
        
        Args:
            excel_path: Path to the Excel file (.xlsx)
            csv_path: Path for the output CSV file
            sheet_name: Name of the sheet to convert (default: active sheet)
        
        Returns:
            Success or error message
        """
        try:
            import openpyxl
        except ImportError:
            return "Error: openpyxl package required. Install with: pip install openpyxl"
        
        try:
            if not os.path.exists(excel_path):
                return f"Error: File not found: {excel_path}"
            
            wb = openpyxl.load_workbook(excel_path)
            
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    return f"Error: Sheet '{sheet_name}' not found"
                ws = wb[sheet_name]
            else:
                ws = wb.active
            
            with open(csv_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                for row in ws.iter_rows(values_only=True):
                    writer.writerow([str(v) if v is not None else '' for v in row])
            
            wb.close()
            return f"Converted {excel_path} to {csv_path}"
        except Exception as e:
            logger.error(f"Failed to convert Excel to CSV: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def csv_to_excel(csv_path: str, excel_path: str, 
                     sheet_name: str = "Sheet1") -> str:
        """
        Convert a CSV file to Excel format.
        Requires openpyxl package.
        
        Args:
            csv_path: Path to the CSV file
            excel_path: Path for the output Excel file (.xlsx)
            sheet_name: Name for the sheet
        
        Returns:
            Success or error message
        """
        try:
            import openpyxl
        except ImportError:
            return "Error: openpyxl package required. Install with: pip install openpyxl"
        
        try:
            if not os.path.exists(csv_path):
                return f"Error: File not found: {csv_path}"
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            with open(csv_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for row in reader:
                    ws.append(row)
            
            wb.save(excel_path)
            wb.close()
            return f"Converted {csv_path} to {excel_path}"
        except Exception as e:
            logger.error(f"Failed to convert CSV to Excel: {e}")
            return f"Error: {e}"

    @mcp_server.tool()
    def excel_to_table(excel_path: str, 
                       sheet_name: Optional[str] = None,
                       max_rows: int = 100) -> str:
        """
        Import an Excel file as a table in the current document.
        Requires openpyxl package.
        
        Args:
            excel_path: Path to the Excel file (.xlsx)
            sheet_name: Name of the sheet to import (default: active sheet)
            max_rows: Maximum rows to import (default 100)
        
        Returns:
            Success or error message
        """
        try:
            import openpyxl
        except ImportError:
            return "Error: openpyxl package required. Install with: pip install openpyxl"
        
        try:
            if not processor.current_document:
                return "Error: No document is open"
            
            if not os.path.exists(excel_path):
                return f"Error: File not found: {excel_path}"
            
            wb = openpyxl.load_workbook(excel_path)
            
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    return f"Error: Sheet '{sheet_name}' not found"
                ws = wb[sheet_name]
            else:
                ws = wb.active
            
            rows = list(ws.iter_rows(max_row=max_rows, values_only=True))
            if not rows:
                wb.close()
                return "Error: Empty sheet"
            
            num_rows = len(rows)
            num_cols = max(len(row) for row in rows)
            
            table = processor.current_document.add_table(
                rows=num_rows, cols=num_cols
            )
            table.style = 'Table Grid'
            
            for r_idx, row_data in enumerate(rows):
                for c_idx, value in enumerate(row_data or []):
                    if c_idx < num_cols:
                        table.cell(r_idx, c_idx).text = str(value) if value else ''
            
            wb.close()
            return f"Added table with {num_rows} rows and {num_cols} columns from Excel"
        except Exception as e:
            logger.error(f"Failed to convert Excel to table: {e}")
            return f"Error: {e}"
